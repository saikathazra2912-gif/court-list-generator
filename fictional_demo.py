"""Standalone fictional-data demonstration. No official template or real records needed."""
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

ROWS = [
    ("AREA ALPHA", "Magistrate Court", "DEMO-001", "Aarav Sen", "REF-101"),
    ("AREA ALPHA", "Magistrate Court", "DEMO-002", "Diya Roy", "REF-102"),
    ("AREA ALPHA", "Magistrate Court", "DEMO-001", "Aarav Sen", "REF-101"),
    ("AREA ALPHA", "Special Court", "DEMO-003", "Ishan Das", "REF-103"),
    ("AREA BETA", "Magistrate Court", "DEMO-004", "Meera Pal", "REF-104"),
    ("AREA BETA", "Sessions Court", "DEMO-005", "Kabir Bose", "REF-105"),
]
COURT_ORDER = {"Magistrate Court": 0, "Special Court": 1, "Sessions Court": 2}

def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Fictional Demo"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "FICTIONAL HEARING LIST — PORTFOLIO DEMO"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=15)
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    ws.append([])
    ws.append(["SL", "Area", "Court", "Case ID", "Participant", "Reference"])
    border = Border(*(Side(style="thin", color="B7C3D0") for _ in range(4)))
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305B86")
        cell.border = border
    ordered = sorted(ROWS, key=lambda r: (r[0], COURT_ORDER.get(r[1], 99), r[3].casefold()))
    serials = defaultdict(dict)
    for area, court, case_id, person, reference in ordered:
        key = (area, court)
        if case_id not in serials[key]:
            serials[key][case_id] = len(serials[key]) + 1
        ws.append([serials[key][case_id], area, court, case_id, person, reference])
        for cell in ws[ws.max_row]:
            cell.border = border
            cell.alignment = Alignment(wrap_text=True)
    ws.conditional_formatting.add(
        f"D4:D{ws.max_row}",
        FormulaRule(
            formula=['AND(D4<>"",COUNTIF($D$4:$D$1000,D4)>1)'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )
    for column, width in {"A": 8, "B": 19, "C": 24, "D": 16, "E": 23, "F": 18}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:F{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:3"
    output = Path(__file__).resolve().parent / "fictional_demo_output.xlsx"
    wb.save(output)
    print(f"Created {output}")

if __name__ == "__main__":
    main()
