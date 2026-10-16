"""Generate a blank, generic portfolio template. No operational data is loaded."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADERS = ["SL", "JID", "Prisoner Name", "Father Name", "FIR/Case No", "W", "B", "Court Name"]

def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Fictional Template"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H2")
    ws["A1"] = "FICTIONAL COURT LIST — BLANK PORTFOLIO TEMPLATE"
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20
    ws.merge_cells("A3:H4")
    ws["A3"] = "GENERIC COURT HEADING — REPLACE WITH FICTIONAL DEMO VALUES"
    ws["A3"].font = Font(bold=True, color="17365D")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="7A8DA5")
    for col, header in enumerate(HEADERS, 1):
        c = ws.cell(5, col, header)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305B86")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in range(6, 12):
            blank = ws.cell(row, col)
            blank.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for letter, width in {"A": 7, "B": 16, "C": 22, "D": 22, "E": 18, "F": 9, "G": 9, "H": 30}.items():
        ws.column_dimensions[letter].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_area = "A1:H11"
    output = Path(__file__).resolve().parent / "fictional_blank_template.xlsx"
    wb.save(output)
    print(f"Created: {output}")

if __name__ == "__main__":
    main()
